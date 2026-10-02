import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import LOGO_PATH

_QUOTE_DASH_MAP = {
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "–": "-",
    "—": "-",
    "…": "...",
    " ": " ",
}

_SECTION_RE = re.compile(r"^\s*(\d+)\.\s*([A-Z][A-Z \-/&]+)\s*$")


def sanitize_text(text: str) -> str:
    """Normalize typographic characters and strip stray markdown for safe rendering."""
    if not text:
        return ""

    cleaned = text
    for src, dst in _QUOTE_DASH_MAP.items():
        cleaned = cleaned.replace(src, dst)

    # Strip markdown emphasis / heading markers while keeping the underlying text.
    cleaned = re.sub(r"^#{1,6}\s*", "", cleaned, flags=re.MULTILINE)
    cleaned = cleaned.replace("**", "").replace("__", "")
    cleaned = re.sub(r"(?<!\w)\*(?!\*)", "", cleaned)
    cleaned = re.sub(r"(?<!\w)_(?!_)", "", cleaned)
    cleaned = cleaned.replace("`", "")

    # Drop characters outside latin-1 so FPDF's core fonts don't choke.
    cleaned = cleaned.encode("latin-1", errors="ignore").decode("latin-1")

    # Collapse excessive blank lines left behind by stripped markdown.
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def parse_terms(text: str) -> list[str]:
    """Split a semicolon-separated terms string into a clean list of strings."""
    if not text:
        return []
    parts = [part.strip().rstrip(".") for part in text.split(";")]
    return [part for part in parts if part]


def _split_sections(text: str):
    """Yield (heading_or_None, body_lines) blocks from sanitized document text."""
    lines = text.splitlines()
    heading = None
    buffer: list[str] = []

    for line in lines:
        match = _SECTION_RE.match(line)
        if match:
            if heading is not None or buffer:
                yield heading, buffer
            heading = f"{match.group(1)}. {match.group(2).strip()}"
            buffer = []
        else:
            buffer.append(line)

    if heading is not None or buffer:
        yield heading, buffer


def _extract_terms(text: str) -> list[str]:
    """Pull the terms list from the TERMS section of a formatted document, if present."""
    for heading, body_lines in _split_sections(text):
        if heading and "TERM" in heading.upper():
            body = "\n".join(body_lines).strip()
            return parse_terms(body)
    return []


def format_docx(text: str, doc_type: str) -> bytes:
    """Render sanitized document text into a formatted .docx file, returned as bytes."""
    import io

    clean_text = sanitize_text(text)
    terms_list = _extract_terms(clean_text)

    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if LOGO_PATH.exists():
        logo_paragraph = document.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_paragraph.add_run().add_picture(str(LOGO_PATH), width=Inches(1.25))

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(doc_type.strip().upper() or "LEGAL DOCUMENT")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.name = "Times New Roman"

    document.add_paragraph()

    for heading, body_lines in _split_sections(clean_text):
        if heading:
            heading_paragraph = document.add_paragraph()
            heading_run = heading_paragraph.add_run(heading)
            heading_run.bold = True
            heading_run.font.size = Pt(12)
            heading_run.font.name = "Times New Roman"

        body = "\n".join(body_lines).strip()
        if body:
            for paragraph_text in body.split("\n"):
                if paragraph_text.strip():
                    document.add_paragraph(paragraph_text.strip())

    if terms_list:
        document.add_paragraph()
        terms_heading = document.add_paragraph()
        terms_heading_run = terms_heading.add_run("TERMS SUMMARY")
        terms_heading_run.bold = True
        terms_heading_run.font.name = "Times New Roman"

        table = document.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        header_cells[0].text = "#"
        header_cells[1].text = "Term"
        for index, term in enumerate(terms_list, start=1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(index)
            row_cells[1].text = term

    section = document.sections[0]
    footer = section.footer
    footer_paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(
        "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
    )
    footer_run.font.size = Pt(9)
    footer_run.font.name = "Times New Roman"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _LegalPDF(FPDF):
    def __init__(self, doc_type: str, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.doc_type = doc_type.strip().upper() or "LEGAL DOCUMENT"
        self.set_auto_page_break(auto=True, margin=25)

    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=10, y=8, w=14)
        self.set_font("Helvetica", "B", 14)
        self.set_xy(0, 12)
        self.cell(0, 10, self.doc_type, align="C")
        self.ln(18)
        self.set_line_width(0.3)
        self.line(10, self.get_y(), self.w - 10, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-18)
        self.set_line_width(0.2)
        self.line(10, self.get_y(), self.w - 10, self.get_y())
        self.set_font("Helvetica", "I", 8)
        self.set_y(-15)
        self.cell(
            0,
            10,
            "LegalEase Inc. | contact@legalease.com | All Rights Reserved.",
            align="C",
        )


def format_pdf(text: str, doc_type: str) -> bytes:
    """Render sanitized document text into a formatted PDF, returned as bytes."""
    clean_text = sanitize_text(text)
    terms_list = _extract_terms(clean_text)

    pdf = _LegalPDF(doc_type)
    pdf.add_page()
    pdf.set_font("Times", "", 12)

    for heading, body_lines in _split_sections(clean_text):
        if heading:
            pdf.set_font("Times", "B", 12)
            pdf.multi_cell(0, 7, heading)
            pdf.set_font("Times", "", 12)
            pdf.ln(1)

        body = "\n".join(body_lines).strip()
        if body:
            for paragraph_text in body.split("\n"):
                paragraph_text = paragraph_text.strip()
                if paragraph_text:
                    pdf.multi_cell(0, 6, paragraph_text)
                    pdf.ln(1)

    if terms_list:
        pdf.ln(2)
        pdf.set_font("Times", "B", 12)
        pdf.multi_cell(0, 7, "TERMS SUMMARY")
        pdf.ln(1)
        pdf.set_font("Times", "", 12)
        for term in terms_list:
            pdf.multi_cell(0, 6, f"- {term}")
            pdf.ln(1)

    return bytes(pdf.output())


def format_html_preview(text: str) -> str:
    """Convert sanitized document text into clean HTML paragraphs/headings for preview."""
    clean_text = sanitize_text(text)
    parts = []

    for heading, body_lines in _split_sections(clean_text):
        if heading:
            parts.append(f"<h3>{html.escape(heading)}</h3>")

        body = "\n".join(body_lines).strip()
        if body:
            for paragraph_text in body.split("\n"):
                paragraph_text = paragraph_text.strip()
                if paragraph_text:
                    parts.append(f"<p>{html.escape(paragraph_text)}</p>")

    return "\n".join(parts)
